"""
Vendor quotation check (M18).

extract_text(raw, kind) — plain text from a PDF (pdfplumber) or XLSX
    (openpyxl).  Raises UnreadableDocument when there is no text to read.

check_quotation(q, text, session) — runs maker_agent in MODE: QUOTE on the
    text, keeps only listings that pass the evidence guards, then sets every
    line's reference, band and recommendation and the quotation's vendor
    score.  The model supplies names, quantities, prices and listings only;
    band, recommendation and score come from pricing.py.  Flushes, does not
    commit.  A Langflow failure leaves the quotation "pending" with
    check_error; it never raises.
"""
from __future__ import annotations

import asyncio
import io
import logging
from datetime import date

from sqlalchemy.orm import Session

from app.models import Quotation, QuotationLine
from app.services import pricing
from app.services.evidence import guard_listings
from app.services.langflow_client import LangflowError
from app.services.matcher import match_item, normalize_name
from app.services.web_check import _MAX_QUOTES, _save_quotes, call_maker, maker_input

logger = logging.getLogger(__name__)

MAX_TEXT_CHARS = 30_000
CHECK_FAILED_MESSAGE = "AI tidak bisa membaca penawaran ini. Coba lagi atau isi manual."


class UnreadableDocument(Exception):
    """The file has no text layer or cannot be opened."""


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def _cell_text(value: object) -> str:
    # openpyxl returns whole-number cells as float; write them without ".0".
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _xlsx_text(raw: bytes) -> str:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — any parse failure means unreadable
        raise UnreadableDocument("File Excel tidak bisa dibuka.") from exc
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


def _pdf_text(raw: bytes) -> str:
    import pdfplumber

    try:
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            return "\n\n".join((page.extract_text() or "") for page in pdf.pages)
    except Exception as exc:  # noqa: BLE001
        raise UnreadableDocument("File PDF tidak bisa dibuka.") from exc


def extract_text(raw: bytes, kind: str) -> str:
    """Return the document text (capped at MAX_TEXT_CHARS)."""
    text = _pdf_text(raw) if kind == "pdf" else _xlsx_text(raw)
    text = text.strip()
    if not text:
        if kind == "pdf":
            raise UnreadableDocument("PDF tidak punya lapisan teks; masukkan manual.")
        raise UnreadableDocument("File Excel kosong.")
    return text[:MAX_TEXT_CHARS]


# ---------------------------------------------------------------------------
# Check
# ---------------------------------------------------------------------------

def _parse_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        return None


def check_quotation(q: Quotation, text: str, session: Session) -> None:
    try:
        maker, results = asyncio.run(call_maker(maker_input("QUOTE", text)))
    except LangflowError:
        logger.exception("maker_agent failed for quotation id=%s", q.id)
        q.status = "pending"
        q.check_error = CHECK_FAILED_MESSAGE
        session.flush()
        return

    q.vendor_name = maker.vendor.name
    q.quote_no = maker.vendor.quote_no
    q.quote_date = _parse_date(maker.vendor.date)
    q.quote_total_idr = maker.quote_total_idr

    score_lines: list[pricing.ScoreLine] = []
    dropped_count = 0
    for line_no, item in enumerate(maker.items, start=1):
        guarded = guard_listings(item.listings, results)
        kept = guarded.kept[:_MAX_QUOTES]
        dropped_count += len(guarded.dropped)

        # Reference: purchase history first, else the median of kept quotes.
        matched = match_item(item.name, session)
        ref: int | None = None
        source = "none"
        history_count = 0
        if matched is not None:
            ref, history_count = pricing.reference_from_history(matched.id, session)
            if ref is not None:
                source = "history"
        if ref is None:
            ref = pricing.reference_from_quotes([l.price_idr for l in kept])
            if ref is not None:
                source = "web"

        diff = pricing.diff_pct(item.unit_price_idr, ref)
        band = pricing.get_price_status(diff)

        offers = [
            {"title": l.title, "url": l.url, "seller": l.seller, "price_idr": l.price_idr}
            for l in kept
        ]
        if offers:
            # Cache them so the to-queue re-check reuses these quotes.
            _save_quotes(normalize_name(item.name), item.search_query[:255], offers, session)

        q.lines.append(
            QuotationLine(
                line_no=line_no,
                name=item.name[:255],
                qty=item.qty,
                unit=item.unit,
                unit_price_idr=item.unit_price_idr,
                line_total_idr=item.line_total_idr,
                item_id=matched.id if matched is not None else None,
                reference_unit_price_idr=ref,
                reference_source=source,
                history_count=history_count or None,
                diff_pct=diff,
                band=band,
                recommendation=pricing.recommend(band, diff),
                search_query=item.search_query[:500],
                evidence=offers,
                dropped=[
                    {"title": d.title, "url": d.url, "price_idr": d.price_idr, "reason": d.reason}
                    for d in guarded.dropped
                ],
            )
        )
        score_lines.append(
            pricing.ScoreLine(band, item.qty, item.unit_price_idr, item.line_total_idr)
        )

    score = pricing.vendor_score(score_lines, maker.quote_total_idr)
    q.vendor_score = score.score
    q.price_points = score.price_points
    q.arithmetic_ok = score.arithmetic_ok
    q.covered_lines = score.covered
    q.total_lines = score.total
    q.dropped_count = dropped_count
    q.status = "checked"
    q.check_error = None
    session.flush()
